#!/usr/bin/env python3
"""Build task BC only. Run with the Codex bundled Python (python-docx, Pillow).

python scripts/build-submission-docs.py --allow-placeholders  # layout draft
python scripts/build-submission-docs.py                      # submission gate
python scripts/build-submission-docs.py --render             # also render QA PNGs

No asset files are created or changed. Missing/ambiguous images fail final builds.
Only 01_经纬_参赛计划书 outputs and its dedicated QA directory are written.
"""
from __future__ import annotations
import argparse
import os
from pypdf import PdfReader
import re
import shutil
import subprocess
import tempfile
from pathlib import Path
from docx import Document
from docx.enum.text import WD_TAB_ALIGNMENT
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image

ROOT = Path(__file__).resolve().parents[1]
RUNTIME = Path.home() / '.cache/codex-runtimes/codex-primary-runtime/dependencies'
STEM = '01_经纬_参赛计划书'
TITLE = '经纬·规则e判'


def font(style, name, size, bold=False):
    style.font.name = name
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor.from_string('25252B')
    fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
    fonts.set(qn('w:eastAsia'), name)
    for key in ('asciiTheme', 'hAnsiTheme', 'eastAsiaTheme', 'cstheme'):
        fonts.attrib.pop(qn('w:' + key), None)


def field(p, instruction, cache=''):
    r = p.add_run()
    begin = OxmlElement('w:fldChar'); begin.set(qn('w:fldCharType'), 'begin')
    text = OxmlElement('w:instrText'); text.set(qn('xml:space'), 'preserve'); text.text = instruction
    sep = OxmlElement('w:fldChar'); sep.set(qn('w:fldCharType'), 'separate')
    value = OxmlElement('w:t'); value.text = cache
    end = OxmlElement('w:fldChar'); end.set(qn('w:fldCharType'), 'end')
    for e in (begin, text, sep, value, end): r._r.append(e)


def inline(p, text):
    for i, part in enumerate(re.split(r'\*\*(.*?)\*\*', text)):
        r = p.add_run(part); r.bold = bool(i % 2)


def table(doc, lines):
    rows = [[x.strip() for x in line.strip().strip('|').split('|')] for line in lines]
    rows = [r for r in rows if not all(re.fullmatch(r':?-+:?', x) for x in r)]
    t = doc.add_table(rows=0, cols=len(rows[0])); t.style = 'Table Grid'
    t.autofit = False
    # Most tables compare values; leave space for descriptive first/last columns.
    n = len(rows[0]); weights = [1] * n
    if n == 3: weights = [1.0, 1.9, 2.2]
    elif n == 5 and '新增资金' in rows[0]: weights = [1.15, .6, .65, .9, 1.9]
    elif n == 5: weights = [1.15, 1, 1, 1, 1]
    elif n == 4: weights = [1.4, 1, 1, 1.5]
    widths = [Cm(16 * w / sum(weights)) for w in weights]
    for col, width in zip(t.columns, widths): col.width = width
    for ri, values in enumerate(rows):
        row = t.add_row()
        trpr = row._tr.get_or_add_trPr()
        no_split = OxmlElement('w:cantSplit'); trpr.append(no_split)
        if ri == 0:
            repeat = OxmlElement('w:tblHeader'); trpr.append(repeat)
        for ci, value in enumerate(values):
            cell = row.cells[ci]; cell.width = widths[ci]
            p = cell.paragraphs[0]; p.paragraph_format.first_line_indent = Cm(0)
            p.paragraph_format.space_after = Pt(4); p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.line_spacing = Pt(12)
            inline(p, value)
            for r in p.runs: r.font.size = Pt(9); r.bold = ri == 0
            if ri == 0:
                fill = OxmlElement('w:shd'); fill.set(qn('w:fill'), 'F1E9EB'); cell._tc.get_or_add_tcPr().append(fill)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument('--allow-placeholders', action='store_true')
    ap.add_argument('--render', action='store_true')
    ap.add_argument('--soffice', type=Path, default=RUNTIME / 'bin/override/soffice')
    args = ap.parse_args()
    source = ROOT / 'docs/经纬_参赛计划书.md'
    output = ROOT / 'exports/submission'; output.mkdir(parents=True, exist_ok=True)
    lines = source.read_text().splitlines()
    missing = []
    assets = {}
    for line in lines:
        match = re.fullmatch(r'!\[(.*?)\]\((.*?)\)', line)
        if match:
            label, pattern = match.groups()
            matches = [p for p in source.parent.glob(pattern) if p.suffix.lower() in {'.png', '.jpg', '.jpeg'}]
            # BB provides PNG plus SVG; use PNG only. Each numbered asset is unique.
            if len(matches) != 1:
                missing.append(f'{label}: {pattern} ({len(matches)} images)')
            else: assets[pattern] = matches[0]
    if missing and not args.allow_placeholders:
        ap.error('素材未齐或编号不唯一，拒绝生成提交版：\n' + '\n'.join(missing))
    if not args.soffice.is_file(): ap.error(f'Bundled soffice missing: {args.soffice}')
    doc = Document()
    sec = doc.sections[0]; sec.page_width = Cm(21); sec.page_height = Cm(29.7)
    sec.top_margin = Cm(2.25); sec.bottom_margin = Cm(2.1)
    sec.left_margin = sec.right_margin = Cm(2.5)
    sec.header_distance = sec.footer_distance = Cm(1.0)
    sec.different_first_page_header_footer = True
    font(doc.styles['Normal'], 'Songti SC', 11)
    normal = doc.styles['Normal'].paragraph_format
    normal.line_spacing = Pt(17); normal.space_after = Pt(7); normal.first_line_indent = Pt(22)
    font(doc.styles['Title'], 'Songti SC', 22, True)
    for style in doc.styles:
        for border in style.element.xpath('.//w:pBdr'):
            border.getparent().remove(border)
    for name, size in [('Heading 1', 16), ('Heading 2', 12.5)]:
        font(doc.styles[name], 'Heiti SC', size, True)
        fmt = doc.styles[name].paragraph_format
        fmt.first_line_indent = Cm(0); fmt.space_before = Pt(12); fmt.space_after = Pt(10)
        fmt.keep_with_next = True
    font(doc.styles['Caption'], 'Songti SC', 10, True)
    cap = doc.styles['Caption'].paragraph_format; cap.first_line_indent = Cm(0); cap.space_after = Pt(9)
    header = sec.header.paragraphs[0]; header.alignment = 2
    header.paragraph_format.first_line_indent = Cm(0)
    header.add_run(TITLE + '  参赛计划书').font.size = Pt(9)
    footer = sec.footer.paragraphs[0]; footer.alignment = 1
    footer.paragraph_format.first_line_indent = Cm(0)
    field(footer, ' PAGE ', '1')
    update = OxmlElement('w:updateFields'); update.set(qn('w:val'), 'true'); doc.settings.element.append(update)
    toc_entries = []
    cover = True; i = 0; body_pages = 1
    while i < len(lines):
        line = lines[i].strip(); i += 1
        if not line: continue
        if line == '<!-- cover-end -->':
            cover = False; doc.add_page_break()
            p = doc.add_paragraph('目录', 'Title'); p.alignment = 1
            headings = [x[3:] for x in lines if x.startswith('## ')]
            for hi, heading in enumerate(headings):
                tp = doc.add_paragraph()
                tp.paragraph_format.first_line_indent = Cm(0)
                tp.paragraph_format.space_after = Pt(14)
                tp.paragraph_format.tab_stops.clear_all()
                tp.paragraph_format.tab_stops.add_tab_stop(Cm(15), WD_TAB_ALIGNMENT.RIGHT)
                if hi == 0:
                    r = tp.add_run()
                    b = OxmlElement('w:fldChar'); b.set(qn('w:fldCharType'), 'begin'); r._r.append(b)
                    ins = OxmlElement('w:instrText'); ins.text = ' TOC \\o "1-1" \\h \\z \\u '; r._r.append(ins)
                    sep = OxmlElement('w:fldChar'); sep.set(qn('w:fldCharType'), 'separate'); r._r.append(sep)
                tp.add_run(heading + '\t')
                number = tp.add_run('0'); toc_entries.append((heading, number))
                if hi == len(headings)-1:
                    r = tp.add_run(); e = OxmlElement('w:fldChar'); e.set(qn('w:fldCharType'), 'end'); r._r.append(e)
            doc.add_page_break(); continue
        if line == '<!-- pagebreak -->':
            bp = doc.add_paragraph(); bp.paragraph_format.line_spacing = Pt(1); bp.paragraph_format.space_after = Pt(0)
            from docx.enum.text import WD_BREAK
            bp.add_run().add_break(WD_BREAK.PAGE); body_pages += 1; continue
        if line.startswith('# '):
            p = doc.add_paragraph(line[2:], 'Title'); p.alignment = 1
            p.paragraph_format.space_before = Pt(95); p.paragraph_format.space_after = Pt(55); continue
        if cover:
            p = doc.add_paragraph(line); p.alignment = 1; p.paragraph_format.first_line_indent = Cm(0)
            p.paragraph_format.space_after = Pt(17); continue
        if line.startswith('## '): doc.add_paragraph(line[3:], 'Heading 1'); continue
        if line.startswith('### '): doc.add_paragraph(line[4:], 'Heading 2'); continue
        if line.startswith('|'):
            block = [line]
            while i < len(lines) and lines[i].strip().startswith('|'): block.append(lines[i].strip()); i += 1
            table(doc, block); continue
        match = re.fullmatch(r'!\[(.*?)\]\((.*?)\)', line)
        if match:
            label, pattern = match.groups(); p = doc.add_paragraph(); p.alignment = 1
            p.paragraph_format.line_spacing = 1.0
            p.paragraph_format.first_line_indent = Cm(0); p.paragraph_format.keep_with_next = True
            if pattern in assets:
                path = assets[pattern]
                with Image.open(path) as im: w, h = im.size
                # OOXML viewport crops leave the BA files and embedded originals intact.
                bottom = 0
                if '/screens/06' in pattern: bottom = round((h-1250)/h*100000)
                elif '/screens/08' in pattern: bottom = round((h-1650)/h*100000)
                visible_h = h * (1 - bottom / 100000)
                max_h = 6.7 if '/screens/' in pattern else 7.2
                width = min(16, max_h * w / visible_h)
                shape = p.add_run().add_picture(str(path), width=Cm(width), height=Cm(width * visible_h / w))
                if bottom:
                    fill = shape._inline.xpath('.//pic:blipFill')[0]
                    crop = OxmlElement('a:srcRect'); crop.set('b', str(bottom)); fill.insert(1, crop)
                pic = p._p.xpath('.//wp:docPr')
                if pic: pic[0].set('descr', label)
            else:
                p.add_run('【素材待补：' + label + '】').font.color.rgb = RGBColor.from_string('963749')
                p.paragraph_format.space_before = Pt(40); p.paragraph_format.space_after = Pt(40)
            continue
        if re.match(r'[图表]\d+ ', line):
            p = doc.add_paragraph(line, 'Caption'); p.alignment = 1
            if line.startswith('表'): p.paragraph_format.keep_with_next = True
            continue
        p = doc.add_paragraph(); inline(p, line)
        if line.startswith('['):
            p.paragraph_format.first_line_indent = Cm(0)
            for r in p.runs: r.font.size = Pt(9)
    doc.core_properties.title = lines[0][2:]; doc.core_properties.author = ''; doc.core_properties.subject = '财富管理服务'
    docx = output / (STEM + '.docx'); pdf = output / (STEM + '.pdf')
    # Keep a real Word TOC field; materialize its cached page numbers from PDF bookmarks.
    # Export into an isolated temporary directory. Publish the matching DOCX and PDF only after conversion succeeds.
    with tempfile.TemporaryDirectory(prefix='jingwei-bc-') as tmp:
        td = Path(tmp); draft = td / docx.name; doc.save(draft)
        config = td / 'fonts.conf'
        config.write_text('<?xml version="1.0"?><!DOCTYPE fontconfig SYSTEM "fonts.dtd"><fontconfig><dir>/System/Library/Fonts</dir><dir>/System/Library/Fonts/Supplemental</dir><cachedir>' + str(td/'font-cache') + '</cachedir></fontconfig>')
        env = dict(os.environ, FONTCONFIG_FILE=str(config))
        previous = None
        for attempt in range(3):
            doc.save(draft)
            built = td / pdf.name
            if built.exists(): built.unlink()
            result = subprocess.run([str(args.soffice), '-env:UserInstallation=' + (td/'profile').as_uri(), '--headless', '--convert-to', 'pdf', '--outdir', str(td), str(draft)], capture_output=True, text=True, timeout=120, env=env)
            if result.returncode or not built.is_file(): raise RuntimeError(result.stdout + result.stderr)
            reader = PdfReader(str(built))
            destinations = {}
            for item in reader.outline:
                if not isinstance(item, list): destinations[item.title] = reader.get_destination_page_number(item) + 1
            pages = [destinations[heading] for heading, _ in toc_entries]
            if pages == previous: break
            previous = pages
            for (_, run), number in zip(toc_entries, pages): run.text = str(number)
        else: raise RuntimeError('目录页码在三次排版后仍不稳定')
        shutil.copyfile(draft, docx); shutil.copyfile(built, pdf)
        if args.render:
            renderer = Path.home()/'.codex/plugins/cache/openai-primary-runtime/documents/26.905.11957/skills/documents/render_docx.py'
            import sys
            subprocess.run([sys.executable, str(renderer), str(docx), '--output_dir', str(output/'qa-bc'), '--emit_pdf'], check=True, env=env)
    print(f'Built {docx}\nBuilt {pdf}\nPlanned content sections: {body_pages}')
    if missing: print('占位草稿，尚不可提交：\n' + '\n'.join(missing))


if __name__ == '__main__': main()
