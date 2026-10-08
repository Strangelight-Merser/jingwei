#!/usr/bin/env python3
"""Build task BH's editable deck. Run again after BA/BB assets arrive.

Bundled Python: ~/.cache/codex-runtimes/codex-primary-runtime/dependencies/python/bin/python3
Pass --require-assets for the submission build; missing assets then stop the build.
Only 04_经纬_路演.* and previews under deck-preview/ are written.
"""
import argparse
import json
import os
import re
import subprocess
import tempfile
import urllib.request
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from PIL import Image
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_TICK_LABEL_POSITION
from pptx.chart.data import CategoryChartData
from pptx.oxml.xmlchemy import OxmlElement

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'exports/submission'
RUNTIME = Path.home() / '.cache/codex-runtimes/codex-primary-runtime/dependencies'
INK, RED, PAPER, MUTED = '25252B', '963749', 'F4F4F6', '6B6B74'
TITLE = '经纬·规则e判——基于公开估值规则与大模型解读的指数基金决策服务方案'


def personal_cover():
    path = OUT / 'personal.json'
    personal = json.loads(path.read_text(encoding='utf-8')) if path.is_file() else {}
    fields = {}
    for key in ('姓名', '学校', '学号', '指导教师', '联系方式'):
        value = personal.get(key, '【待填】')
        if not isinstance(value, str):
            raise ValueError(f'个人信息字段必须为字符串：{key}')
        fields[key] = f'{key}：{value or "【待填】"}'
    return '\n'.join(('财富管理服务（兼顾青年群体服务）',
                      '  '.join(fields[k] for k in ('姓名', '学校', '学号')),
                      '  '.join(fields[k] for k in ('指导教师', '联系方式')),
                      '2026-10-08'))



def text(slide, value, x, y, w, h, size=24, color=INK, bold=False, font='PingFang SC'):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = 0
    frame.margin_top = frame.margin_bottom = 0
    for i, line in enumerate(value.split('\n')):
        p = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        p.text = line
        p.font.name = font
        p.font.size = Pt(size)
        p.font.bold = bold
        p.font.color.rgb = RGBColor.from_string(color)
        p.line_spacing = 1.3
        p.space_after = Pt(14)
    return box


def asset(kind, number, keyword):
    folder = OUT / 'assets' / kind
    files = sorted(folder.glob('*')) if folder.exists() else []
    files = [p for p in files if p.suffix.lower() in ('.png', '.jpg', '.jpeg')]
    numbered = [p for p in files if re.match(rf'^0?{number}(?:\D|$)', p.stem)]
    matches = numbered or [p for p in files if keyword in p.stem]
    if len(matches) > 1:
        raise ValueError(f'素材匹配不唯一：{kind}/{number} {keyword}: {matches}')
    return matches[0] if matches else None


def picture(slide, path, x, y, w, h, crop_top=0, crop_bottom=0, rounded=True, crop_left=0, crop_right=0):
    with Image.open(path) as im:
        ratio = im.width*(1-crop_left-crop_right) / (im.height * (1-crop_top-crop_bottom))
    width, height = (w,w/ratio) if ratio >= w/h else (h*ratio,h)
    shape=slide.shapes.add_picture(str(path),Inches(x+(w-width)/2),Inches(y+(h-height)/2),width=Inches(width),height=Inches(height))
    shape.crop_top,shape.crop_bottom=crop_top,crop_bottom
    shape.crop_left,shape.crop_right=crop_left,crop_right
    if rounded:
        geom=shape._element.spPr.find('{http://schemas.openxmlformats.org/drawingml/2006/main}prstGeom')
        geom.set('prst','roundRect')
        guide=OxmlElement('a:gd'); guide.set('name','adj'); guide.set('fmla',f'val {int((12/72)/min(width,height)*100000)}')
        geom[0].append(guide)
        effects=OxmlElement('a:effectLst'); shadow=OxmlElement('a:outerShdw')
        for key,value in {'blurRad':'152400','dist':'50800','dir':'5400000','algn':'ctr','rotWithShape':'0'}.items(): shadow.set(key,value)
        color=OxmlElement('a:srgbClr'); color.set('val',INK)
        alpha=OxmlElement('a:alpha'); alpha.set('val','16000')
        color.append(alpha); shadow.append(color); effects.append(shadow); shape._element.spPr.append(effects)
    return shape


def card(slide,x,heading,fact,detail):
    sh=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(x),Inches(2),Inches(3.85),Inches(4.35))
    sh.fill.solid(); sh.fill.fore_color.rgb=RGBColor.from_string(PAPER); sh.line.fill.background(); sh.adjustments[0]=.035
    sh._element.spPr.append(OxmlElement('a:effectLst'))
    text(slide,heading,x+.3,2.35,3.25,.6,23,bold=True)
    text(slide,fact,x+.3,3.45,3.25,1.25,28,RED,True)
    text(slide,detail,x+.3,5.05,3.25,1.05,18)


def icon(slide,x,kind):
    def line(a,b,c,d):
        sh=slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(a),Inches(b),Inches(c),Inches(d))
        sh.line.color.rgb=RGBColor.from_string(RED); sh.line.width=Pt(2)
    if kind==0:
        sh=slide.shapes.add_shape(MSO_SHAPE.OVAL,Inches(x),Inches(3.3),Inches(.8),Inches(.8))
        sh.fill.background(); sh.line.color.rgb=RGBColor.from_string(RED); sh.line.width=Pt(2)
        line(x+.4,3.5,x+.4,3.9); line(x+.2,3.7,x+.6,3.7)
    elif kind==1:
        for k in range(3): line(x+k*.3,4.1,x+k*.3,3.7-k*.2)
        line(x-.1,4.2,x+.8,4.2)
    else:
        line(x,4.1,x+.3,3.8); line(x+.3,3.8,x+.55,3.95); line(x+.55,3.95,x+.85,3.4); line(x+.55,3.3,x+.55,4.2)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--require-assets', action='store_true')
    parser.add_argument('--api', default='http://localhost:4411')
    args = parser.parse_args()
    with urllib.request.urlopen(args.api + '/publication/judgments', timeout=20) as response:
        indexes = json.load(response)['indexes']
    j = next(row for row in indexes if row['index_code'] == '000300')
    with urllib.request.urlopen(args.api + '/publication/rule-outcomes?index=000300', timeout=20) as response:
        outcomes = json.load(response)
    n = len(indexes)
    if n != 11:
        raise ValueError('覆盖指数已变化，请先更新路演内容')
    b = j['boundaries']
    stable = (ROOT / 'docs/规则稳健性.md').read_text()
    if '49/50' not in stable or '17 次完整排序' not in stable:
        raise ValueError('稳健性结果已变化，请先更新路演内容')
    # title, concise copy, visual, speaker notes, source caption
    slides = [
      ('经纬让每次基金判断都有可解释的依据。', TITLE + '\n财富管理服务（兼顾青年群体服务）\n姓名、学校、学号、指导教师、联系方式：【待填】\n2026-10-08', None,
       '新钱怎么安排，手里的基金怎么看，什么变化会让判断改变？经纬把公开规则、历史结果和AI解读接在一起，给客户和客户经理同一份清楚的依据。', ''),
      ('客户需要知道怎么安排资金，也需要知道何时改变判断。', '新钱怎么安排？\n已有基金怎么看？\n什么变化会让判断改变？', None,
       '客户面对的困难是判断依据难以连起来。基金名称不同，可能跟踪同一指数。市场变化后，原来的说法为什么改变，也需要解释。', ''),
      ('公开规则给出判断，大模型把依据讲清楚。', '数据有来源\n规则可复算\n解释跟随所选指数', ('figures',1,'总体架构'),
       '经纬用公开估值数据计算规则结果，提供新增资金和持仓两种说法。大模型读取规则片段解释原因，输出还要通过程序校验。', '产品总体架构'),
      ('今日判断分别回答新钱与已有持仓的问题。', f"沪深300：{j['pe_ttm']:.2f}倍，第{j['percentile']:.1f}百分位\n{j['judgment']['new_money']['title']}\n已有持仓：{j['judgment']['held']['title']}", ('screens',1,'首页沪深300'),
       f"当前沪深300是{j['pe_ttm']:.2f}倍、第{j['percentile']:.1f}百分位，处于{j['judgment']['label']}。首页区分新增资金与已有持仓。拖动曲线可以回看历史位置。", f"当前软件公开判断；市场数据截至{j['as_of']}"),
      ('缓冲与连续确认让判断在边界附近更稳定。', f"四区间边界：30 / 70 / 90\n离开当前区间设5个百分点缓冲\n日数据5个数据日，周数据2个周读数", ('figures',3,'规则'),
       f"规则先看十年PE分位，再用五个百分点缓冲和连续确认减少反复改口。沪深300当前参考边界为{b['low']:.2f}、{b['high']:.2f}、{b['extreme']:.2f}倍。边界随历史窗口变化。", '现行v2；当前参考边界来自软件公开判断'),
      ('历史回放提供证据，收益结论保留样本边界。', f"沪深300v2历史回算：{j['stability']['changes']}次改判\n稳健性：有样本区间排序49/50次成立\n四档齐全18次，其中17次完整排序成立", ('screens',6,'收益'),
       '判断变化页给出含分红的收益回放和股债性价比第二视角，动作仍按PE规则。稳健性覆盖五个指数、十组参数，49次有样本区间排序成立。只有18次四档齐全，17次完整排序成立。历史起点持有期重叠，回算不代表当时发布或未来收益。', 'docs/规则稳健性.md；月初取样，含分红、未扣基金费用'),
      ('11个指数明确区分官方日数据与第三方周数据。', '中证官网每日估值：6个指数\n蛋卷基金每周估值：5个指数\n全收益与ERP：5个指数', ('screens',2,'全部指数'),
       '六个中证指数使用中证官网日估值。创业板指和四个境外指数使用蛋卷基金周估值，页面明确标注第三方。日频和周频分别确认。收益回放和ERP覆盖五个有对应序列的指数。历史不足十年的指数按已有历史标注。', 'docs/指数覆盖.md；全收益和ERP不覆盖所有11个指数'),
      ('持仓体检按指数汇总，识别名字不同的重复配置。', '导入后核对基金与金额\n查看方向、区间和重复配置\n导出一页体检报告', ('screens',8,'体检结果'),
       '这页画面为标注清楚的虚构持仓示例。截图在本机识别，核对后保存。普通跟踪基金按对应指数给判断，指数增强、主动混合、货币等不套用规则。体检帮助客户发现不同名字对应同一个指数，报告便于留存和沟通。', '虚构持仓示例；规则覆盖11个指数的普通跟踪基金'),
      ('问经纬通过校验后展示与规则一致的解释。', '规则决定判断\n模型解释原因\n校验不通过时显示规则原文', ('figures',4,'问经纬'),
       '问经纬跟随所选指数取证，模型生成解释后，程序检查立场、数字、日期、指数、来源、确认长度和预测用语。每条依据可展开看原文。演示不临时发起付费请求。', '规则判别与大模型解读分工'),
      ('工行基金频道与客户经理可以共用一份判断依据。', '基金频道嵌入规则卡与JSON\n客户经理生成一页说明\n试点测客户理解与解释耗时', ('screens',12,'规则卡'),
       '拟在工银手机银行基金频道接入规则卡，客户经理用同一数据生成一页说明。两种渠道保持数据日期、来源和改判条件一致。当前是落地方案，尚未接入银行生产系统。试点先测理解和解释耗时，再评价运营价值。', '工行落地方案；当前提供规则卡、嵌入代码和JSON地址'),
      ('上线审校与数据标注共同限定服务边界。', '接续风险测评、适当性匹配与留痕\n第三方数据标明来源与频率\n持仓本机保存，模型解释受规则校验约束', None,
       '银行上线需要审校动作措辞，并接续正式测评、风险匹配、揭示与留痕。第三方与官方数据不能直接互换，历史排序也不保证未来。规则结果不构成个别投资建议。密钥由用户自备，模型调用需要费用授权。', '第三方对照见docs/指数覆盖.md；历史收益不保证未来'),
      ('桌面与网页已形成从判断到报告的完整流程。', '今日判断、改判记录与收益回放\n持仓体检、问经纬与机构服务\n下一步：工行试点与Windows真机验收', ('figures',2,'用户流程'),
       '当前已实现桌面和网页的完整流程，支持11个指数。规则有独立复算和稳健性检验。具体交付包仍需最终质检，Windows原生安装、OCR和PDF等仍需真机验收。试点围绕客户理解、解释耗时和实际使用反馈推进。', '实施现状与后续试点；不把交付包待验收项写成已通过'),
      ('让每次判断讲得清，让每次改变有依据。', '经纬·规则e判\n基于公开估值规则与大模型解读的\n指数基金决策服务方案', None,
       '经纬把可复算的规则判断带到客户的持仓和客户经理的沟通里。让每次判断讲得清，让每次改变有依据。谢谢。', ''),
    ]
    slides.insert(5, ('沪深300历史上低估值区间的三年年化均值更高。',
        '\n'.join(f"{row['label']}  {row['three_year']['mean']:.1f}%" for row in outcomes['bands']), None,
        '按逐日已确认的区间分组，沪深300四档之后的三年年化均值为9.6%、6.3%、3.7%、负5.1%。这是含分红指数收益，未扣基金费用。逐日样本的持有期重叠，一年期没有同样排序，不能把这张表读成收益承诺。下一页用月初取样的参数检验和ERP补充观察。',
        f"软件收益回放；截至{outcomes['returns_as_of']}；含分红、未扣基金费用，逐日样本有重叠"))
    resolved = [asset(*s[2]) if s[2] else None for s in slides]
    missing = [f'{i+1}: {s[2]}' for i,s in enumerate(slides) if s[2] and resolved[i] is None]
    if args.require_assets and missing:
        raise SystemExit('BA/BB素材尚未齐全：\n'+'\n'.join(missing))
    copies={
      2:'公开数据有来源\n估值规则可复算\n解释跟随所选指数',
      3:f"沪深300 PE {j['pe_ttm']:.2f} 倍\n第 {j['percentile']:.1f} 百分位\n新钱与持仓分别判断",
      4:'四区间：30 / 70 / 90\n离开区间缓冲 5 点\n日 5 次、周 2 次确认',
      6:f"历史回算 {j['stability']['changes']} 次改判\n区间排序成立 49/50 次\n完整排序成立 17/18 次",
      7:'官方日估值覆盖 6 指数\n第三方周估值覆盖 5 指数\n收益与 ERP 覆盖 5 指数',
      8:'导入后核对基金与金额\n识别同向重复配置\n导出一页体检报告',
      9:'规则决定判断\n模型解释原因\n校验失败展示规则原文',
      10:'基金频道共用规则卡\n客户经理生成一页说明\n试点测理解与解释耗时',
      12:'桌面与网页流程已形成\n判断、持仓、解释与报告\n后续试点及 Windows 验收'}
    home=asset('screens',1,'首页沪深300'); mobile=asset('screens',16,'手机首页'); holdings=asset('screens',8,'体检结果')
    if not all((home,mobile,holdings)): raise SystemExit('BH需要01、08、16截图')
    prs=Presentation(); prs.slide_width,prs.slide_height=Inches(13.333333),Inches(7.5)
    for i,(title,copy,visual,notes,source) in enumerate(slides):
        slide=prs.slides.add_slide(prs.slide_layouts[6]); slide.background.fill.solid()
        dark=i in (0,13)
        slide.background.fill.fore_color.rgb=RGBColor.from_string('1F1F24' if dark else 'FFFFFF')
        if dark:
            text(slide,'经纬·规则e判',.6,1.35,6.5,.9,44,'FFFFFF',True)
            subtitle='基于公开估值规则与大模型解读的\n指数基金决策服务方案' if i==0 else title
            text(slide,subtitle,.6,2.65,5.7,1.4,20,'C9C9D0')
            if i==0: text(slide,title,.6,4.55,5.5,.9,18,'C9C9D0')
            metadata = text(slide,personal_cover(),.6,5.8,6.4,1.15,11,'C9C9D0')
            for paragraph in metadata.text_frame.paragraphs:
                paragraph.space_after = Pt(0)
            picture(slide,home,7.25,1.4,7,4.8,.29,.065,crop_left=.11,crop_right=.11)
        else:
            text(slide,title,.6,.6,12.13,1,28,bold=True)
            if i==1:
                for k,(heading,detail) in enumerate([('新钱怎么投','临时追加需要明确依据'),('已有基金怎么看','名字不同也可能投向相同'),('什么时候会改判','市场变化需要可回看的理由')]):
                    x=.6+k*4.14; card(slide,x,heading,'',detail); icon(slide,x+.35,k)
            elif i==5:
                data=CategoryChartData(number_format='0.0'); data.categories=['偏低','中间','偏高','高位']; values=[9.6,6.3,3.7,-5.1]
                actual=[round(row['three_year']['mean'],1) for row in outcomes['bands']]
                if actual!=values: raise ValueError(f'收益数据已变化：{actual}')
                data.add_series('三年年化均值',values)
                chart=slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED,Inches(.75),Inches(1.95),Inches(8.25),Inches(4.6),data).chart
                chart.has_legend=False; chart.has_title=False; chart.chart_style=2
                chart.series[0].invert_if_negative=False
                plot=chart.plots[0]; plot.gap_width=75; plot.has_data_labels=True
                labels=plot.data_labels; labels.position=XL_LABEL_POSITION.OUTSIDE_END; labels.number_format='0.0\\%;-0.0\\%'; labels.font.size=Pt(20); labels.font.name='PingFang SC'
                # Explicit point-level sign handling preserves negative columns in LibreOffice.
                for point,color in zip(chart.series[0].points,['3F9A6E','6B7C99','C98A3A','C9536C']):
                    invert=OxmlElement('c:invertIfNegative'); invert.set('val','0'); point._ser.get_or_add_dPt_for_point(point._idx).insert(1,invert)
                    point.format.fill.solid(); point.format.fill.fore_color.rgb=RGBColor.from_string(color); point.format.line.fill.background()
                chart.value_axis.minimum_scale=-8; chart.value_axis.maximum_scale=12; chart.value_axis.major_unit=4
                chart.category_axis.tick_label_position=XL_TICK_LABEL_POSITION.LOW
                chart.value_axis.tick_labels.number_format='0"%"'; chart.value_axis.tick_labels.font.size=Pt(12)
                chart.category_axis.tick_labels.font.size=Pt(18); chart.category_axis.tick_labels.font.name='PingFang SC'
                text(slide,'沪深300，v2 规则\n三年年化均值\n逐日样本，含分红\n未扣基金费用\n\n样本持有期重叠\n不代表未来',9.45,2.35,3.2,3.9,18,MUTED)
            elif i==11:
                card(slide,.6,'合规审核','上线前审校','接续风险测评、适当性匹配与留痕')
                card(slide,4.74,'第三方数据','64%–81%','五个A股指数同日周读数的 v2 判断一致率')
                card(slide,8.88,'模型解释','规则原文保留','解释受规则校验约束，未通过则展示规则原文')
                text(slide,'科创50仅20.31%，64周样本；境外指数尚无官方PE对照，不跨来源替换。',.9,6.52,11.5,.4,12,MUTED)
            elif i==12:
                text(slide,copies[i],.6,2.05,4.2,3.6,18)
                picture(slide,mobile,5.05,1.8,2.18,4.85)
                picture(slide,holdings,7.55,1.8,5.18,4.85,0,.51)
                text(slide,'手机首页',5.05,6.65,2.18,.3,12,MUTED)
                text(slide,'持仓体检 · 虚构持仓示例',7.55,6.65,5.18,.3,12,MUTED)
            else:
                text(slide,copies[i],.6,2.1,4.1,3.8,18)
                if resolved[i]:
                    top,bottom=0,0
                    if visual[0]=='screens':
                        if i==3: top,bottom=.29,.065
                        elif i==6: top,bottom=.33,.32
                        elif i==8: bottom=.51
                        elif i==10: bottom=.43
                    picture(slide,resolved[i],4.9,1.8,7.83,4.9,top,bottom,rounded=visual[0]=='screens',crop_left=.1 if visual[0]=='screens' and i not in (6,8) else 0,crop_right=.1 if visual[0]=='screens' and i not in (6,8) else 0)
                else: text(slide,'素材待接入',4.9,3,7.8,1,24,MUTED)
            text(slide,source,.6,7.03,11.6,.25,9,MUTED)
        text(slide,f'{i+1:02d}',12.23,7.03,.5,.25,9,'C9C9D0' if dark else MUTED)
        slide.notes_slide.notes_text_frame.text=notes
    OUT.mkdir(parents=True,exist_ok=True)
    target = OUT / '04_经纬_路演.pptx'
    prs.save(target)
    soffice = RUNTIME / 'bin/override/soffice'
    with tempfile.TemporaryDirectory(prefix='jingwei-bd-office-') as profile:
        config = Path(profile) / 'fonts.conf'
        config.write_text('<?xml version="1.0"?><!DOCTYPE fontconfig SYSTEM "fonts.dtd"><fontconfig><dir>/System/Library/Fonts</dir><dir>/System/Library/Fonts/Supplemental</dir><dir>/Library/Fonts</dir><cachedir>' + profile + '/font-cache</cachedir></fontconfig>')
        env = dict(os.environ, FONTCONFIG_FILE=str(config))
        subprocess.run([str(soffice),f'-env:UserInstallation={Path(profile).as_uri()}',
                        '--headless','--convert-to','pdf','--outdir',str(OUT),str(target)],check=True,env=env)
    preview = OUT / 'deck-preview'
    preview.mkdir(exist_ok=True)
    subprocess.run([str(RUNTIME/'bin/override/pdftoppm'),'-scale-to','1600','-png',
                    str(target.with_suffix('.pdf')),str(preview/'slide')],check=True)
    print(f'已生成{len(slides)}页：{target}')
    print('结构稿，等待素材：\n'+'\n'.join(missing) if missing else '全部BA/BB素材已接入。')

if __name__ == '__main__':
    main()
